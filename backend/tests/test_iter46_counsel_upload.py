"""Iteration 46 — Counsel middleware refactor + full-notice upload endpoint.

Modules under test:
  - backend/utils/readonly_admin.py  (ReadonlyEnforcementMiddleware)
  - backend/routers/legal.py          (POST /api/legal/docs/{slug}/upload)
"""
import io
import os
import shutil
from pathlib import Path

import pytest
import requests
from dotenv import dotenv_values

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
if not base_url:
    raise RuntimeError("REACT_APP_BACKEND_URL missing")
BASE_URL = base_url.rstrip("/")

COUNSEL = ("counsel@birthright.live", "counsel-review-2026")
ADMIN = ("admin@birthright.live", "birthright2026")

SLUG = "03-cookie-notice"
SOURCE = Path("/app/backend/legal_docs/03-cookie-notice.md")
SNAPSHOT = Path("/tmp/iter46_03-cookie-notice.md.bak")
DOCX_SOURCE = Path("/app/backend/legal_docs/03-cookie-notice.docx")
DOCX_SNAPSHOT = Path("/tmp/iter46_03-cookie-notice.docx.bak")


def _login(email, password):
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=30)
    if r.status_code != 200:
        pytest.fail(f"login failed for {email}: {r.status_code} {r.text[:300]}")
    body = r.json()
    return body["token"], body["user"]


@pytest.fixture(scope="module")
def counsel():
    token, user = _login(*COUNSEL)
    assert user["role"] == "readonly_admin", f"expected readonly_admin, got {user['role']}"
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {token}"})
    return s


@pytest.fixture(scope="module")
def admin():
    token, user = _login(*ADMIN)
    assert user["role"] == "admin"
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {token}"})
    return s


@pytest.fixture(scope="module", autouse=True)
def restore_sources():
    """Snapshot the legal source (+ built docx) before mutating; restore after."""
    shutil.copy2(SOURCE, SNAPSHOT)
    if DOCX_SOURCE.exists():
        shutil.copy2(DOCX_SOURCE, DOCX_SNAPSHOT)
    yield
    shutil.copy2(SNAPSHOT, SOURCE)
    if DOCX_SNAPSHOT.exists():
        shutil.copy2(DOCX_SNAPSHOT, DOCX_SOURCE)


def _is_readonly_block(resp):
    if resp.status_code != 403:
        return False
    try:
        return resp.json().get("readonly") is True
    except Exception:
        return False


# ---------------- Middleware behaviour ----------------
class TestCounselMiddleware:
    def test_admin_mutation_blocked(self, counsel):
        r = counsel.post(f"{BASE_URL}/api/admin/settings/counsel/rotate", json={}, timeout=30)
        assert r.status_code == 403, f"expected 403, got {r.status_code}: {r.text[:300]}"
        body = r.json()
        assert body.get("readonly") is True
        assert "off-limits for the counsel account" in body.get("detail", "")

    def test_admin_legal_prefix_allowlisted(self, counsel):
        r = counsel.post(f"{BASE_URL}/api/admin/legal/anything", json={}, timeout=30)
        assert not _is_readonly_block(r), f"admin/legal/* should be allow-listed, got {r.status_code} {r.text[:200]}"

    def test_legal_comment_write_allowed(self, counsel):
        r = counsel.post(
            f"{BASE_URL}/api/legal/comments/{SLUG}",
            json={"body": "TEST_iter46 counsel redline", "kind": "comment"},
            timeout=30,
        )
        assert r.status_code == 200, f"counsel legal comment failed: {r.status_code} {r.text[:300]}"
        data = r.json()
        assert data["source_slug"] == SLUG
        assert data["author_role"] == "readonly_admin"
        assert "_id" not in data
        cid = data["id"]
        # verify persistence
        lst = counsel.get(f"{BASE_URL}/api/legal/comments/{SLUG}", timeout=30)
        assert lst.status_code == 200
        assert any(c["id"] == cid for c in lst.json())

    def test_non_admin_mutation_not_readonly_blocked(self, counsel):
        """Counsel should reach normal end-user routes (no readonly 403)."""
        for path, payload in [
            ("/api/cart/items", {"product_id": "nope", "quantity": 1}),
            ("/api/auth/update-profile", {"first_name": "Counsel"}),
        ]:
            r = counsel.post(f"{BASE_URL}{path}", json=payload, timeout=30)
            assert not _is_readonly_block(r), f"{path} readonly-blocked: {r.text[:200]}"

    def test_admin_role_unaffected(self, admin):
        r = admin.get(f"{BASE_URL}/api/legal/comments/{SLUG}", timeout=30)
        assert r.status_code == 200


# ---------------- Upload endpoint ----------------
def _md_bytes(marker):
    return f"# Cookie Notice\n\nTEST_iter46 {marker}\n\nBody paragraph.\n".encode("utf-8")


def _docx_bytes():
    from docx import Document
    doc = Document()
    doc.add_heading("Cookie Notice", level=1)
    doc.add_paragraph("TEST_iter46 docx upload body.")
    doc.add_heading("Section Two", level=2)
    doc.add_paragraph("Second paragraph.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestFullUpload:
    def test_counsel_md_upload(self, counsel):
        r = counsel.post(
            f"{BASE_URL}/api/legal/docs/{SLUG}/upload",
            files={"file": ("replacement.md", _md_bytes("md-counsel"), "text/markdown")},
            timeout=180,
        )
        assert r.status_code == 200, f"{r.status_code} {r.text[:400]}"
        d = r.json()
        for key in ("bytes_written", "body_hash", "rebuilt_docx", "message"):
            assert key in d, f"missing {key} in {d}"
        assert d["source_slug"] == SLUG
        assert isinstance(d["bytes_written"], int) and d["bytes_written"] > 0
        assert isinstance(d["body_hash"], str) and len(d["body_hash"]) > 0
        # persistence: source file rewritten
        assert "TEST_iter46 md-counsel" in SOURCE.read_text(encoding="utf-8")

    def test_counsel_docx_upload(self, counsel):
        r = counsel.post(
            f"{BASE_URL}/api/legal/docs/{SLUG}/upload",
            files={"file": ("replacement.docx", _docx_bytes(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            timeout=180,
        )
        assert r.status_code == 200, f"{r.status_code} {r.text[:400]}"
        d = r.json()
        for key in ("bytes_written", "body_hash", "rebuilt_docx", "message"):
            assert key in d
        text = SOURCE.read_text(encoding="utf-8")
        assert "TEST_iter46 docx upload body." in text
        assert "# Cookie Notice" in text, "heading not converted to markdown"

    def test_admin_md_upload(self, admin):
        r = admin.post(
            f"{BASE_URL}/api/legal/docs/{SLUG}/upload",
            files={"file": ("replacement.md", _md_bytes("md-admin"), "text/markdown")},
            timeout=180,
        )
        assert r.status_code == 200, f"{r.status_code} {r.text[:400]}"
        assert "TEST_iter46 md-admin" in SOURCE.read_text(encoding="utf-8")

    def test_missing_file_400(self, counsel):
        r = counsel.post(f"{BASE_URL}/api/legal/docs/{SLUG}/upload",
                         files={"notfile": ("x.md", b"hello")}, timeout=60)
        assert r.status_code == 400, f"{r.status_code} {r.text[:300]}"

    def test_unknown_slug_404(self, counsel):
        r = counsel.post(f"{BASE_URL}/api/legal/docs/TEST_nope-slug/upload",
                         files={"file": ("x.md", b"hello")}, timeout=60)
        assert r.status_code == 404, f"{r.status_code} {r.text[:300]}"

    def test_bad_extension_400(self, counsel):
        r = counsel.post(f"{BASE_URL}/api/legal/docs/{SLUG}/upload",
                         files={"file": ("x.pdf", b"%PDF-1.4 junk")}, timeout=60)
        assert r.status_code == 400, f"{r.status_code} {r.text[:300]}"
        assert "Unsupported" in r.text or "unsupported" in r.text

    def test_empty_file_400(self, counsel):
        r = counsel.post(f"{BASE_URL}/api/legal/docs/{SLUG}/upload",
                         files={"file": ("x.md", b"")}, timeout=60)
        assert r.status_code == 400, f"{r.status_code} {r.text[:300]}"

    def test_anonymous_rejected(self):
        r = requests.post(f"{BASE_URL}/api/legal/docs/{SLUG}/upload",
                          files={"file": ("x.md", b"hello")}, timeout=60)
        assert r.status_code in (401, 403), f"{r.status_code} {r.text[:200]}"
